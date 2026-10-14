import io
from datetime import datetime
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

EXCEL_REPORT_COLUMNS = [
    "Candidate Name",
    "Email ID",
    "Contact Number",
    "Total Experience",
    "Role Applied For",
    "JD Match Score"
]

def generate_candidate_excel_report(candidates_data: list[dict], default_role: str = "General") -> bytes:
    """
    Generate standard professional .xlsx report adhering strictly to ATS requirements:
    1. Candidate Name
    2. Email ID
    3. Contact Number
    4. Total Experience
    5. Role Applied For
    6. JD Match Score (MUST ALWAYS be the last column)

    Styled with bold header, frozen top row, column filters, and auto-adjusted widths.
    """
    rows = []
    for c in candidates_data:
        name = c.get("Candidate Name") or c.get("full_name") or "Not Found"
        email = c.get("Email ID") or c.get("Email") or c.get("email") or "Not Found"
        phone = c.get("Contact Number") or c.get("Phone") or c.get("phone") or "Not Found"
        exp = c.get("Total Experience") or c.get("total_experience") or "Not Available"
        role = c.get("Role Applied For") or c.get("job_title") or default_role or "Not Specified"
        
        raw_score = c.get("JD Match Score") if "JD Match Score" in c else c.get("Score (%)", c.get("match_score", 0.0))
        try:
            score = round(float(raw_score), 2)
        except (ValueError, TypeError):
            score = 0.0

        rows.append({
            "Candidate Name": str(name),
            "Email ID": str(email),
            "Contact Number": str(phone),
            "Total Experience": str(exp),
            "Role Applied For": str(role),
            "JD Match Score": score
        })

    # Sort descending by JD Match Score by default
    rows = sorted(rows, key=lambda x: x["JD Match Score"], reverse=True)

    df = pd.DataFrame(rows, columns=EXCEL_REPORT_COLUMNS)

    output = io.BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Screened Candidates")
        worksheet = writer.sheets["Screened Candidates"]

        # Freeze the top header row
        worksheet.freeze_panes = "A2"

        # Enable auto-filter across all columns
        worksheet.auto_filter.ref = worksheet.dimensions

        # Styling definitions
        header_fill = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        data_font = Font(name="Calibri", size=11)
        data_alignment = Alignment(vertical="center")
        center_alignment = Alignment(horizontal="center", vertical="center")
        score_alignment = Alignment(horizontal="right", vertical="center")

        thin_border = Border(
            left=Side(style='thin', color="E5E7EB"),
            right=Side(style='thin', color="E5E7EB"),
            top=Side(style='thin', color="E5E7EB"),
            bottom=Side(style='thin', color="E5E7EB")
        )

        # Style header row
        for col_idx in range(1, len(EXCEL_REPORT_COLUMNS) + 1):
            cell = worksheet.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_alignment

        worksheet.row_dimensions[1].height = 28

        # Style data rows
        for row_idx in range(2, len(rows) + 2):
            worksheet.row_dimensions[row_idx].height = 22
            for col_idx in range(1, len(EXCEL_REPORT_COLUMNS) + 1):
                cell = worksheet.cell(row=row_idx, column=col_idx)
                cell.font = data_font
                cell.border = thin_border
                col_name = EXCEL_REPORT_COLUMNS[col_idx - 1]

                if col_name in ["Contact Number", "Total Experience"]:
                    cell.alignment = center_alignment
                elif col_name == "JD Match Score":
                    cell.alignment = score_alignment
                    cell.number_format = '0.0'
                else:
                    cell.alignment = data_alignment

        # Auto-adjust column widths
        for col in worksheet.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val = str(cell.value or '')
                max_len = max(max_len, len(val))
            worksheet.column_dimensions[col_letter].width = max(max_len + 5, 16)

    output.seek(0)
    return output.getvalue()

def generate_failed_resumes_report(failed_data: list[dict]) -> bytes:
    """Generate an Excel report listing failed/unprocessed resumes and reasons."""
    clean_failed = []
    for item in failed_data:
        clean_failed.append({
            "Resume Filename": item.get("Resume Name", "Unknown"),
            "Status": item.get("Status", "Failed"),
            "Error Reason": item.get("Reason", "Unknown error")
        })

    df = pd.DataFrame(clean_failed)
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Failed Extractions")
        worksheet = writer.sheets["Failed Extractions"]
        worksheet.freeze_panes = "A2"
        for col in worksheet.columns:
            col_letter = get_column_letter(col[0].column)
            worksheet.column_dimensions[col_letter].width = 30
    output.seek(0)
    return output.getvalue()

def get_report_filename(prefix: str = "ATS_Candidate_Screening") -> str:
    """Generate standardized filename: ATS_Candidate_Screening_YYYYMMDD.xlsx"""
    date_str = datetime.now().strftime("%Y%m%d")
    return f"{prefix}_{date_str}.xlsx"
