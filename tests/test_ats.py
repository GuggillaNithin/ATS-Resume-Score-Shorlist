import os
import unittest
import tempfile
import shutil
import io
import openpyxl
import warnings
import logging

os.environ["HF_HUB_DISABLE_SYMLINKS_WARNING"] = "1"
os.environ["TOKENIZERS_PARALLELISM"] = "false"
warnings.filterwarnings("ignore")
logging.getLogger("transformers").setLevel(logging.ERROR)
logging.getLogger("sentence_transformers").setLevel(logging.ERROR)
logging.getLogger("huggingface_hub").setLevel(logging.ERROR)


from modules.database import init_db, get_db
from modules.resume_parser import (
    extract_resume_text,
    extract_candidate_name,
    extract_email,
    extract_phone,
    extract_experience,
    extract_skills
)
from modules.jd_matcher import clean_text, calculate_jd_match_score, batch_calculate_match_scores
from modules.excel_export import generate_candidate_excel_report, EXCEL_REPORT_COLUMNS
from modules.candidate_service import (
    get_or_create_candidate,
    add_resume_record,
    search_candidates,
    update_candidate_details,
    delete_candidate
)
from modules.screening_service import (
    create_requisition,
    process_single_resume,
    rescreen_candidates_for_new_jd
)

class TestATSResumeScreening(unittest.TestCase):
    def setUp(self):
        self.test_dir = tempfile.mkdtemp()
        self.db_path = os.path.join(self.test_dir, "test_ats.db")
        init_db(self.db_path)

    def tearDown(self):
        shutil.rmtree(self.test_dir, ignore_errors=True)

    def test_text_cleaning_and_matching(self):
        # Verify clean_text preserves keywords and strips stopwords
        cleaned = clean_text("Looking for a Senior Python Developer with AWS and Docker experience")
        self.assertIn("senior", cleaned)
        self.assertIn("python", cleaned)
        self.assertIn("aws", cleaned)
        self.assertNotIn("with", cleaned)
        self.assertNotIn("for", cleaned)

        # Verify score calculation returns float 0-100
        score = calculate_jd_match_score(
            "Python machine learning scikit-learn",
            "Expert Python developer specializing in machine learning algorithms and scikit-learn."
        )
        self.assertGreater(score, 50.0)
        self.assertLessEqual(score, 100.0)

    def test_candidate_extraction_heuristics(self):
        sample_resume_text = """
        John Doe
        Software Engineer
        Email: john.doe@example.com | Phone: +91 9876543210
        Summary: 5+ years of experience in Python development.
        
        Experience:
        Senior Developer | XYZ Corp | Jan 2021 - Present
        Junior Developer | ABC Inc | Jan 2019 - Dec 2020
        
        Skills:
        Python, Docker, AWS, React, PostgreSQL
        """

        name = extract_candidate_name(sample_resume_text, "John_Doe_Resume.pdf")
        self.assertEqual(name, "John Doe")

        email = extract_email(sample_resume_text)
        self.assertEqual(email, "john.doe@example.com")

        phone = extract_phone(sample_resume_text)
        self.assertEqual(phone, "9876543210")

        exp = extract_experience(sample_resume_text)
        self.assertTrue("Years" in exp or "5" in exp)

        skills = extract_skills(sample_resume_text)
        self.assertIn("Python", skills)
        self.assertIn("AWS", skills)
        self.assertIn("Docker", skills)

    def test_scanned_pdf_detection(self):
        empty_pdf_bytes = b"%PDF-1.4\n1 0 obj\n<<>>\nendobj\ntrailer\n<<>>\n%%EOF"
        text, err = extract_resume_text(empty_pdf_bytes, "scanned_doc.pdf")
        self.assertTrue(err is not None)
        self.assertTrue("requires OCR" in err or "Failed" in err or "Scanned" in err)

    def test_candidate_deduplication(self):
        # 1. Create first candidate
        cand1, is_new1, _ = get_or_create_candidate(
            full_name="Alice Smith",
            email="alice.smith@example.com",
            phone="9123456789",
            total_experience="3 Years",
            current_role="Data Scientist",
            skills=["Python", "SQL"],
            file_hash="hash_abc_123",
            db_path=self.db_path
        )
        self.assertTrue(is_new1)

        # 2. Upload same candidate with same email -> deduplicate
        cand2, is_new2, note2 = get_or_create_candidate(
            full_name="Alice S.",
            email="ALICE.SMITH@EXAMPLE.COM", # Case-insensitive
            phone="9123456789",
            total_experience="4 Years",
            current_role="Senior Data Scientist",
            skills=["Python", "SQL", "Pandas"],
            file_hash="hash_def_456",
            db_path=self.db_path
        )
        self.assertFalse(is_new2)
        self.assertEqual(cand1["candidate_id"], cand2["candidate_id"])
        self.assertIn("email", note2.lower())

        # 3. Candidate with same file hash -> deduplicate
        cand3, is_new3, note3 = get_or_create_candidate(
            full_name="Unknown Name",
            email="other_email@example.com",
            phone="9998887776",
            total_experience=None,
            current_role=None,
            skills=[],
            file_hash="hash_abc_123",
            db_path=self.db_path
        )
        # First candidate resume record had hash_abc_123 if added
        add_resume_record(cand1["candidate_id"], "alice.pdf", "/fake/path", "hash_abc_123", "text", "Completed", db_path=self.db_path)
        cand4, is_new4, note4 = get_or_create_candidate(
            full_name="Unknown",
            email="new_email@example.com",
            phone=None,
            total_experience=None,
            current_role=None,
            skills=[],
            file_hash="hash_abc_123",
            db_path=self.db_path
        )
        self.assertFalse(is_new4)
        self.assertEqual(cand1["candidate_id"], cand4["candidate_id"])

    def test_excel_export_format(self):
        candidates = [
            {
                "Candidate Name": "Alice Smith",
                "Email ID": "alice@example.com",
                "Contact Number": "9123456789",
                "Total Experience": "4.5 Years",
                "Role Applied For": "AI Engineer",
                "JD Match Score": 88.5
            },
            {
                "Candidate Name": "Bob Jones",
                "Email ID": "bob@example.com",
                "Contact Number": "9876543210",
                "Total Experience": "2 Years",
                "Role Applied For": "AI Engineer",
                "JD Match Score": 72.1
            }
        ]

        excel_bytes = generate_candidate_excel_report(candidates, default_role="AI Engineer")
        self.assertIsNotNone(excel_bytes)
        self.assertGreater(len(excel_bytes), 1000)

        # Inspect workbook structure using openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
        sheet = wb["Screened Candidates"]

        # Check headers match EXACTLY
        header_vals = [cell.value for cell in sheet[1]]
        self.assertEqual(header_vals, EXCEL_REPORT_COLUMNS)
        # Check that JD Match Score is the last column
        self.assertEqual(header_vals[-1], "JD Match Score")

        # Check descending order by score
        first_row_score = sheet.cell(row=2, column=6).value
        second_row_score = sheet.cell(row=3, column=6).value
        self.assertGreaterEqual(float(first_row_score), float(second_row_score))

    def test_end_to_end_screening_and_new_jd_search(self):
        # 1. Create a Requisition
        req = create_requisition("Backend Engineer", "Python, FastAPI, PostgreSQL, Docker, AWS", db_path=self.db_path)
        self.assertIsNotNone(req["requisition_id"])

        # 2. Add Candidate via valid DOCX in-memory creation
        import docx
        doc = docx.Document()
        doc.add_paragraph("Carol Danvers")
        doc.add_paragraph("Email: carol@example.com | Phone: 9988776655")
        doc.add_paragraph("4+ years experience in Python, FastAPI, and Docker.")
        doc.add_paragraph("Senior Python Engineer")
        doc_buf = io.BytesIO()
        doc.save(doc_buf)
        valid_docx_bytes = doc_buf.getvalue()

        cand_res = process_single_resume(
            file_bytes=valid_docx_bytes,
            filename="Carol_Danvers_Resume.docx",
            jd_text="Python, FastAPI, PostgreSQL, Docker, AWS",
            job_title="Backend Engineer",
            requisition_id=req["requisition_id"],
            batch_id="batch_test_01",
            threshold=50.0,
            db_path=self.db_path
        )
        self.assertEqual(cand_res["status"], "Completed")

        self.assertEqual(cand_res["candidate_name"], "Carol Danvers")
        self.assertEqual(cand_res["email"], "carol@example.com")
        self.assertGreater(cand_res["score"], 0.0)

        # 3. Search Candidate in Database
        found_cands, count = search_candidates(search_query="carol", db_path=self.db_path)
        self.assertEqual(count, 1)
        self.assertEqual(found_cands[0]["full_name"], "Carol Danvers")

        # 4. Rescreen existing candidate against a brand new JD
        new_req, rescreened = rescreen_candidates_for_new_jd(
            new_job_title="DevOps Specialist",
            new_jd_text="AWS cloud architecture, Terraform, Kubernetes, CI/CD",
            threshold=40.0,
            save_to_database=True,
            db_path=self.db_path
        )
        self.assertEqual(len(rescreened), 1)
        self.assertEqual(rescreened[0]["Candidate Name"], "Carol Danvers")
        self.assertEqual(rescreened[0]["Role Applied For"], "DevOps Specialist")

        # 5. Delete candidate and ensure clean removal
        cand_id = cand_res["candidate_id"]
        deleted = delete_candidate(cand_id, db_path=self.db_path)
        self.assertTrue(deleted)
        _, count_after = search_candidates(search_query="carol", db_path=self.db_path)
        self.assertEqual(count_after, 0)

if __name__ == "__main__":

    unittest.main()
